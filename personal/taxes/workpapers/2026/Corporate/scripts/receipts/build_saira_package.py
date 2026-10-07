"""Build the 2025-26 HST filing package for Saira: folder of source documents + workpaper."""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from organize_receipts import optimize

TAXES = Path('/Users/umarfarooqaslam/Documents/Taxes')
SCRATCH = Path(__file__).resolve().parents[2] / 'parsed_data' / 'receipt_work'
VEHICLE = Path('/Users/umarfarooqaslam/obsidian/obsidian-me/personal/business-vehicle/docs/lease')
OUT = TAXES / '2026/Corporate/HST Filing 2025-26 - for Saira'
WORKPAPERS = Path(__file__).resolve().parents[4]
WS26 = WORKPAPERS / '2026/Corporate/parsed_data'
WS25 = WORKPAPERS / '2025/Corporate/parsed_data'
LEDGER = WS26 / 'transactions_master.csv'
DIRECTOR = TAXES / '2026/Corporate/Director-funded expenses'
INVOICES = TAXES / '2026/Corporate/Debit Account/Invoices'
PREV_FEE = 'Accounting fee — previous accountant (T2 2025)'

F = Font(name='Arial', size=10)
FB = Font(name='Arial', size=10, bold=True)
FH = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FT = Font(name='Arial', size=14, bold=True)
FIN = Font(name='Arial', size=10, color='0000FF')
FLINK = Font(name='Arial', size=10, color='008000')
HEAD_FILL = PatternFill('solid', fgColor='1F3864')
KEY_FILL = PatternFill('solid', fgColor='FFFF00')
TOTAL_FILL = PatternFill('solid', fgColor='D9E1F2')
THIN = Side(style='thin', color='999999')
MONEY = '#,##0.00;(#,##0.00);-'
REFUNDED = {}


# ---------- documents ----------
def copy_file(src, dst):
    """Copy a document into the package, losslessly compressed."""
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_bytes(optimize(Path(src)))


def key(date, desc, amount):
    return (str(date)[:10], str(desc), round(float(amount), 2))


def copy_docs(write=True):
    """Copy every source document into the package; return the receipt name for each line."""
    if write and OUT.exists():
        raise SystemExit(f'{OUT} exists; move it aside first, or pass --workpaper-only')
    put = copy_file if write else (lambda src, dst: None)
    chq = OUT / '02 Bank statements - BMO chequing'
    cc = OUT / '03 Card statements - BMO Mastercard'
    ls = OUT / '05 Vehicle lease - Mercedes GLC 43'
    pr = OUT / '06 Prior HST return 2024-25'
    d25 = TAXES / '2025/Corporate/Debit Account/Statements'
    for n in ['8-August 29, 2025.pdf', '9-September 29, 2025.pdf', '10-October 31, 2025.pdf',
              '11-November 28, 2025.pdf', '12-December 31, 2025.pdf']:
        put(d25 / n, chq / f'2025-{n}')
    d26 = TAXES / '2026/Corporate/Debit Account/Statements'
    for i, n in enumerate(['January 30', 'February 27', 'March 31', 'April 30', 'May 29',
                           'June 30', 'July 31', 'August 31'], 1):
        put(d26 / f'{n}, 2026.pdf', chq / f'2026-{i:02d}-{n}, 2026.pdf')
    c25 = TAXES / '2025/Corporate/Credit Card/Statements'
    for n in ['08-August 28, 2025.pdf', '09-September 28, 2025.pdf', '10-October 28, 2025.pdf',
              '11-November 28, 2025.pdf', '12-December 28, 2025.pdf']:
        put(c25 / n, cc / f'2025-{n}')
    c26 = TAXES / '2026/Corporate/Credit Card'
    for i, n in enumerate(['January', 'February', 'March', 'April', 'May', 'June', 'July',
                           'August'], 1):
        put(c26 / f'{n} 28, 2026.pdf', cc / f'2026-{i:02d}-{n} 28, 2026.pdf')
    for n in ['lease-agreement-detailed.pdf', 'lease-pricing-worksheet.pdf',
              'payment-receipts.pdf']:
        put(VEHICLE / n, ls / n)
    put(TAXES / '2025/Corporate/HST Filing/15248132 Canada Inc. HST.pdf',
        pr / 'NETFILE confirmation 837076 - 2024-08-15 to 2025-08-14.pdf')
    t2 = (TAXES / '2025/Corporate/accountant files/final'
          / "1. T2 Yr '25 - Client Copy - 15248132 CANADA INC (20251231) - Umar Farooq Aslam.pdf")
    put(t2, OUT / '07 T2 2025 - as filed' / 'T2 2025 - 15248132 Canada Inc. - client copy.pdf')
    for src in sorted(INVOICES.iterdir()):
        if src.is_file() and not src.name.startswith('.'):
            put(src, OUT / '04c Invoices - chequing payments' / src.name)
    for src in sorted(DIRECTOR.iterdir()):
        if src.is_file() and not src.name.startswith('.'):
            put(src, OUT / '05b Director-funded - Bell and Rogers' / src.name)

    names = {}
    inv25 = TAXES / '2025/Corporate/Credit Card/Invoices by Month'
    for _, r in pd.read_csv(WS25 / 'receipts_2025_hst_period_index.csv').iterrows():
        folder, fname = r.receipt.split('/', 1)
        dst = f'{folder} - {fname}'
        stem = fname.rsplit('-invoice', 1)[0] if '-invoice' in fname else None
        pages = sorted((inv25 / folder).glob(f'{stem}-invoice*')) if stem else []
        for page in pages or [inv25 / r.receipt]:
            put(page, OUT / '04 Receipts - Aug 15 to Dec 31 2025' / f'{folder} - {page.name}')
        names[key(r.date, r.description, r.amount)] = f'04/{dst}'
    r26 = TAXES / '2026/Corporate/Credit Card/receipts'
    for _, r in pd.read_csv(WS26 / 'receipts_2026_index.csv').iterrows():
        if r.receipt.startswith('09-Sep'):
            continue
        put(r26 / r.receipt, OUT / '04b Receipts - 2026' / r.receipt)
        if r.status in ('matched', 'personal charge on corporate card'):
            names.setdefault(key(r.date, r.description, r.amount), f'04b/{r.receipt}')
    return names


# ---------- workbook helpers ----------
def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row, i, c)
        cell.font, cell.fill = FH, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.freeze_panes = ws.cell(row + 1, 1)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_body(ws, first, last, ncol, money_cols=()):
    for r in range(first, last + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(r, c)
            if cell.font != FIN and cell.font != FLINK:
                cell.font = F
            if c in money_cols:
                cell.number_format = MONEY


def total_row(ws, row, label_col, label, sums):
    ws.cell(row, label_col, label).font = FB
    for col, formula in sums.items():
        cell = ws.cell(row, col, formula)
        cell.font, cell.number_format = FB, MONEY
    for c in range(1, ws.max_column + 1):
        ws.cell(row, c).fill = TOTAL_FILL
        ws.cell(row, c).border = Border(top=THIN)


# ---------- sheets ----------
def sheet_sales(wb, p):
    ws = wb.create_sheet('Sales')
    ws['A1'], ws['A1'].font = 'Line 101 / 103 — MIR deposits, Aug 15 2025 to Aug 14 2026', FT
    ws['A2'] = ('Counted by date received, the method used on the 2024-25 return. Every '
                'deposit is on the chequing statement named in column F.')
    header(ws, 4, ['Date received', 'Bank description', 'Gross deposit (incl. 13% HST)',
                   'HST (13/113)', 'Net sales', 'Chequing statement'])
    rev = p[(p.category == 'Revenue') & p.description.str.contains('MIR')].sort_values('date')
    r = 5
    for _, x in rev.iterrows():
        ws.cell(r, 1, x.date.date()).number_format = 'yyyy-mm-dd'
        ws.cell(r, 2, 'MIRSERVICESPAY/PAY — MIR Services')
        c = ws.cell(r, 3, float(x.gross))
        c.font = FIN
        ws.cell(r, 4, f'=ROUND(C{r}*13/113,2)')
        ws.cell(r, 5, f'=C{r}-D{r}')
        ws.cell(r, 6, statement_for(x.date))
        r += 1
    last = r - 1
    style_body(ws, 5, last, 6, (3, 4, 5))
    total_row(ws, r, 2, f'Total — {last - 4} deposits',
              {3: f'=SUM(C5:C{last})', 4: f'=SUM(D5:D{last})', 5: f'=SUM(E5:E{last})'})
    ws.cell(r + 2, 1, 'Not included:').font = FB
    ws.cell(r + 3, 1, 'Aug 21, 2026 MIR deposit, 8,542.80 gross — pays for Aug 3–14 work but '
                      'was received after Aug 14. Next period by date received.')
    ws.cell(r + 4, 1, 'Oct 10, 2025 CRA credit 4,960.50 and May 28, 2026 CRA credit 4,644.90 — '
                      'tax refunds, not sales.')
    widths(ws, [14, 36, 22, 14, 14, 30])
    return ws, r


def statement_for(d):
    months = ['January 30', 'February 27', 'March 31', 'April 30', 'May 29', 'June 30',
              'July 31', 'August 31']
    if d.year == 2025:
        return {8: 'August 29, 2025', 9: 'September 29, 2025', 10: 'October 31, 2025',
                11: 'November 28, 2025', 12: 'December 31, 2025'}[d.month]
    return f'{months[d.month - 1]}, 2026'


REASON = {
    'Meals & entertainment': '50% of HST (ETA s.236)',
    'Vehicle — lease': '80% business use × s.67.3 cap (ETA s.235), see Lease cap',
    'Vehicle & fuel': '80% business use of the Mercedes, see Lease cap',
    'Vehicle — maintenance': '80% business use of the Mercedes, see Lease cap',
    'Vehicle — insurance': 'Exempt supply, no HST',
    'Insurance — home office': 'Exempt supply, no HST',
    'Rent — home office': 'Residential rent, exempt',
    'Subcontractor': 'Offshore, services outside Canada; no HST, no self-assessment '
                     '(corp is exclusively commercial)',
    'Bank & transfer fees': 'Financial service, exempt',
    'Bank fees (HST-exempt)': 'Financial service, exempt',
    'Fines & penalties': 'Not a supply',
    'Government fees': 'No HST on Corporations Canada fees',
    'Personal — owed by Umar': 'Personal charge on the corporate card (hospital); repaid by '
                               'Umar through his account. Not claimed',
    'Cash back rebate': 'Card rebate, no HST',
}


CAR_COSTS = ('Vehicle & fuel', 'Vehicle — maintenance')
OWN_CAR = "Umar's own car before the lease; mileage paid, so no ITC"
BUSINESS_USE_REF = "'Lease cap'!$B$22"


def sheet_itcs(wb, p, receipts, lease_ratio_ref):
    ws = wb.create_sheet('ITCs')
    ws['A1'], ws['A1'].font = 'Line 106 — every business purchase in the period', FT
    ws['A2'] = ('Column H is the HST inside the price (13/113) when a Canadian supplier '
                'charged it. Column E is the amount paid, HST included. Column I is the '
                'claimable share. Blue = inputs you can change.')
    cols = ['Date', 'Paid from', 'Supplier / description', 'Category', 'Amount paid',
            'Canadian HST charged?', 'Statement', 'HST in price', 'ITC %', 'ITC claimed',
            'Why not 100%', 'Receipt']
    header(ws, 4, cols)
    keep = p[(p.source != 'debit') | p.category.isin(
        ['Vehicle — lease', 'Vehicle — insurance', 'Professional fees',
         'Software & subscriptions', 'Bank & transfer fees', 'Subcontractor'])]
    keep = keep[keep.direction != 'in'].sort_values(['date', 'source'])
    src_label = {'card': 'BMO Mastercard', 'debit': 'BMO chequing',
                 'personal': 'Umar personal card (reimbursed)'}
    r = 5
    for _, x in keep.iterrows():
        signed = -x.amount if x.direction == 'refund' else x.amount
        charged = 'Yes' if x.hst_in_price != 0 else 'No'
        ws.cell(r, 1, x.date.date()).number_format = 'yyyy-mm-dd'
        ws.cell(r, 2, src_label[x.source])
        ws.cell(r, 3, PREV_FEE if 'ejaz' in x.description.lower() else x.description)
        ws.cell(r, 4, x.category)
        ws.cell(r, 5, float(signed))
        ws.cell(r, 6, charged).font = FIN
        ws.cell(r, 7, x.ref if x.source != 'debit' else statement_for(x.date))
        ws.cell(r, 8, f'=IF(F{r}="Yes",ROUND(E{r}*13/113,2),0)')
        own_car = x.category in CAR_COSTS and charged == 'Yes' and x.itc_pct == 0
        if x.category == 'Vehicle — lease':
            ws.cell(r, 9, f'={lease_ratio_ref}').font = FLINK
        elif x.category in CAR_COSTS and not own_car:
            ws.cell(r, 9, f'={BUSINESS_USE_REF}').font = FLINK
        else:
            ws.cell(r, 9, float(x.itc_pct) if charged == 'Yes' else 0).font = FIN
        ws.cell(r, 10, f'=ROUND(H{r}*I{r},2)')
        why = OWN_CAR if own_car else REASON.get(x.category, '')
        if charged == 'No' and not why:
            why = 'Non-resident supplier, no Canadian HST' if x.supplier == 'non-resident' \
                else 'No HST in price'
        ws.cell(r, 11, why)
        ws.cell(r, 12, receipt_status(x, receipts))
        r += 1
    last = r - 1
    style_body(ws, 5, last, 12, (5, 8, 10))
    for rr in range(5, last + 1):
        ws.cell(rr, 9).number_format = '0.00%'
    total_row(ws, r, 3, 'Totals (line 106 is ITC claimed, column J)',
              {5: f'=SUM(E5:E{last})', 8: f'=SUM(H5:H{last})', 10: f'=SUM(J5:J{last})'})
    ws.auto_filter.ref = f'A4:L{last}'
    widths(ws, [12, 22, 40, 24, 12, 11, 22, 11, 8, 11, 42, 44])
    return ws, r


def refunded_purchases(p):
    """Keys of card purchases fully reversed by a later credit of the same amount."""
    out = {}
    refunds = p[(p.source == 'card') & (p.direction == 'refund')]
    buys = p[(p.source == 'card') & (p.direction == 'out')]
    for rf in refunds.itertuples():
        stem = rf.description.split()[0][:4].upper()
        hit = buys[(buys.amount == rf.amount) & (buys.date <= rf.date)
                   & buys.description.str.upper().str.startswith(stem)]
        if len(hit):
            b = hit.iloc[-1]
            out[key(b.date, b.description, b.amount)] = rf.date.date().isoformat()
    return out


def receipt_status(x, receipts):
    k = key(x.date, x.description, x.amount)
    if k in receipts:
        return receipts[k]
    if x.direction == 'refund':
        return 'Credit on card statement'
    if k in REFUNDED:
        return f'Refunded in full ({REFUNDED[k]})'
    if x.source == 'debit' and 'vercel' in x.description.lower():
        return '04b/02-Feb/vercel receipt.pdf'
    if x.source == 'personal' and x.category == 'Telecom':
        return '05b/' + ('Bell payment history' if 'Bell' in x.description
                         else 'Rogers bill history')
    if x.source == 'personal' and 'Mercedes' in x.description:
        return '04b/04-April/02-Apr-MercedesBenz-353.94-VISA1680-invoice-p1.jpg (+p2, cardslip)'
    if x.itc_claimable == 0:
        return '—'
    if x.source == 'debit' and x.category == 'Professional fees':
        if 'ejaz' in x.description.lower():
            return 'Invoice on file with Umar (not included)'
        return '04c/09-Feb-CPACM-invoice-2041-423.75.pdf'
    if x.source == 'debit':
        return ('Lease: 05 Vehicle lease (monthly debit)' if x.category == 'Vehicle — lease'
                else 'NEEDED — invoice')
    return 'NEEDED — receipt' if x.amount >= 30 else 'recommended (under $30)'


def sheet_instalments(wb):
    ws = wb.create_sheet('Instalments')
    ws['A1'], ws['A1'].font = 'Line 110 — instalments paid toward 2025-26', FT
    header(ws, 3, ['Date', 'Bank description', 'Amount', 'Chequing statement'])
    rows = [('2025-10-20', 'October 31, 2025'), ('2026-01-20', 'January 30, 2026'),
            ('2026-04-20', 'April 30, 2026'), ('2026-07-20', 'July 31, 2026')]
    for r, (d, st) in enumerate(rows, 4):
        ws.cell(r, 1, pd.Timestamp(d).date()).number_format = 'yyyy-mm-dd'
        ws.cell(r, 2, 'CANADA TXD/DIM — GST prepayment')
        ws.cell(r, 3, 6500.0).font = FIN
        ws.cell(r, 4, st)
    style_body(ws, 4, 7, 4, (3,))
    total_row(ws, 8, 2, 'Total (line 110)', {3: '=SUM(C4:C7)'})
    notes = [
        ('Not included: Sep 15, 2025 payment of 5,800.00. It settled the 2024-25 balance '
        '(35,121.50 owed less 34,282.00 prepaid = 839.50), overpaying by 4,960.50, which CRA '
        'refunded Oct 10, 2025.'),
        ('Please confirm in My Business Account that CRA credited all four payments to the '
        '2025-08-15 to 2026-08-14 period.')]
    for r, n in enumerate(notes, 10):
        ws.cell(r, 1, n).font = F
    widths(ws, [14, 36, 14, 22])
    return ws


def sheet_lease(wb):
    ws = wb.create_sheet('Lease cap')
    title = 'Vehicle lease — 80% business use, capped under s.67.3 / s.235'
    ws['A1'], ws['A1'].font = title, FT
    rows = [
        ('Vehicle', '2026 Mercedes-Benz GLC 43 AMG 4MATIC, VIN W1NKM8HB1TF465706', None),
        ('Lessee', '15248132 Canada Inc. (co-lessee Umar Farooq Aslam)', None),
        ('Lease signed', '2025-09-30, 36 months, Mercedes-Benz Financial', None),
        ('Monthly payment before tax', 1548.68, 'money'),
        ('Monthly HST', 201.33, 'money'),
        ('Down payment (cash)', 4424.78, 'money'),
        ('HST on first payment + down payment, paid at signing', 776.55, 'money'),
        ('Security deposit', 0, 'money'),
        ('CCA / lease ceiling for leases entered in 2025', 38000, 'money'),
        ('List price used: base + factory options (pricing worksheet "total sale price")',
         99500, 'money'),
        ('Alternative: total list incl. freight and dealer items', 109929.95, 'money'),
    ]
    for r, (k, v, kind) in enumerate(rows, 3):
        ws.cell(r, 1, k).font = F
        c = ws.cell(r, 2, v)
        c.font = FIN if kind else F
        if kind:
            c.number_format = MONEY
    ws['B12'].fill = KEY_FILL
    ws['B12'].comment = Comment('Source: lease pricing worksheet (05 Vehicle lease). '
                                'Change to B13 if you use the total list price.', 'Umar')
    ws['A15'], ws['A15'].font = 'Ratio = ceiling ÷ (0.85 × list price)', FB
    ws['B15'] = '=ROUND(B11/(0.85*B12),4)'
    ws['B15'].number_format, ws['B15'].font = '0.00%', FB
    ws['A16'] = 'Monthly $1,100 limit, for comparison (ratio binds when lower)'
    ws['B16'] = '=ROUND(1100/B6,4)'
    ws['B16'].number_format = '0.00%'
    ws['A18'], ws['A18'].font = 'HST paid on the lease inside the period', FB
    ws['A19'], ws['B19'] = 'At signing, Sep 27 + Sep 30, 2025 (on Mastercard)', '=B9'
    ws['A20'], ws['B20'] = 'Monthly debits Oct 30, 2025 to Jul 30, 2026 (10)', '=10*B7'
    ws['A21'], ws['B21'] = 'Total HST paid', '=B19+B20'
    ws['A22'], ws['A22'].font = 'Business use of the Mercedes (lease, fuel, maintenance)', FB
    ws['B22'] = 0.8
    ws['B22'].number_format, ws['B22'].font, ws['B22'].fill = '0.00%', FIN, KEY_FILL
    ws['B22'].comment = Comment("Umar's estimate; see Questions 2. The ITCs sheet uses this "
                                'for Mercedes fuel and maintenance.', 'Umar')
    ws['A23'], ws['A23'].font = 'Share of lease HST claimed = business use × cap (ITCs sheet)', FB
    ws['B23'] = '=ROUND(B22*B15,4)'
    ws['B23'].number_format, ws['B23'].font = '0.00%', FB
    ws['A24'], ws['B24'] = 'ITC claimed', '=ROUND(B21*B23,2)'
    for r in (19, 20, 21, 24):
        ws.cell(r, 2).number_format = MONEY
        ws.cell(r, 1).font = ws.cell(r, 2).font = FB if r in (21, 24) else F
    ws['A26'] = ('The Aug 31, 2026 debit falls in the next period. No mileage log exists for '
                 'the Mercedes yet.')
    ws['A27'] = ('The s.67.3 cap on the income tax deduction is a separate T2 matter and is not '
                 'affected by this sheet.')
    widths(ws, [70, 22])
    return 'ROUND(\'Lease cap\'!$B$23,4)'


def sheet_return(wb, sales_total_row, itc_total_row):
    ws = wb['Return']
    ws['A1'], ws['A1'].font = 'GST/HST return — 15248132 Canada Inc. (XYZ Bytes)', FT
    ws['A2'] = 'BN 795920958 RT0001 · annual filer · period 2025-08-15 to 2026-08-14'
    ws['A3'] = 'Due 2026-11-14 (Saturday, so 2026-11-16). Prepared 2026-09-23 for Saira.'
    header(ws, 5, ['Line', 'Description', 'Amount', 'Source'])
    s, i = sales_total_row, itc_total_row
    lines = [
        ('101', 'Sales and other revenue', f"=Sales!E{s}", 'Sales sheet'),
        ('103', 'GST/HST collected or collectible', f"=Sales!D{s}", 'Sales sheet'),
        ('104', 'Adjustments to be added', 0, 's.173 vehicle benefit — see Questions'),
        ('105', 'Total GST/HST and adjustments', '=C7+C8', ''),
        ('106', 'ITCs', f"=ITCs!J{i}", 'ITCs sheet'),
        ('107', 'Adjustments to be deducted', 0, ''),
        ('108', 'Total ITCs and adjustments', '=C10+C11', ''),
        ('109', 'Net tax', '=C9-C12', ''),
        ('110', 'Instalments', '=Instalments!C8', 'Instalments sheet'),
        ('114', 'Refund claimed', '=MAX(C14-C13,0)', ''),
        ('115', 'Balance owing', '=MAX(C13-C14,0)', ''),
    ]
    for r, (ln, d, v, src) in enumerate(lines, 6):
        ws.cell(r, 1, ln)
        ws.cell(r, 2, d)
        c = ws.cell(r, 3, v)
        c.number_format = MONEY
        if isinstance(v, str) and '!' in v:
            c.font = FLINK
        ws.cell(r, 4, src)
    style_body(ws, 6, 16, 4)
    for r in range(6, 17):
        ws.cell(r, 3).number_format = MONEY
        if ws.cell(r, 1).value in ('109', '114'):
            for c in range(1, 5):
                ws.cell(r, c).font = FB
                ws.cell(r, c).fill = TOTAL_FILL
    ws['A18'], ws['A18'].font = 'Last year (filed 2025-09-24, conf. 837076)', FB
    prior = [('101', 300168.74), ('103', 39021.93), ('106', 3900.43), ('109', 35121.50)]
    for r, (ln, v) in enumerate(prior, 19):
        ws.cell(r, 1, ln).font = F
        c = ws.cell(r, 3, v)
        c.font, c.number_format = FIN, MONEY
    widths(ws, [8, 40, 16, 44])


def sheet_readme(wb):
    ws = wb.active
    ws.title = 'Read me'
    text = [
        ('HST 2025-26 — filing package for Saira', FT),
        (('15248132 Canada Inc. (XYZ Bytes) · BN 795920958 RT0001 · period 2025-08-15 to '
         '2026-08-14'), FB),
        ('', F),
        ('What is in the folder', FB),
        ('01  This workpaper. Return lines, every sale, every purchase with its HST treatment.',
         F),
        (('02  BMO chequing statements, Aug 2025 to Aug 2026 (13). Every month-end balance '
         'ties to the transaction workbook.'), F),
        ('03  BMO Mastercard statements, Aug 2025 to Aug 2026 (13).', F),
        (('04  Receipts for card purchases Aug 15 to Dec 31, 2025, named by month folder. The '
         'Receipt column on the ITCs sheet points to each file.'), F),
        (('04b Receipts for 2026 card purchases by statement month, and the Mercedes service '
         'paid on Umar\'s Visa (04-April).'), F),
        ('05b Bell payment history and Rogers bill history (phone and internet Umar pays).', F),
        ('05  Mercedes lease agreement, pricing worksheet and signing receipts.', F),
        ('06  The 2024-25 NETFILE confirmation you filed (837076).', F),
        ('04c Invoices for purchases paid from the chequing account.', F),
        ('07  T2 2025 as filed May 18, 2026 (client copy).', F),
        ('', F),
        ('How the numbers were built', FB),
        (('Sales are MIR deposits counted by date received, as on last year\'s return. '
         'Purchases come from the two statements plus items Umar pays personally and the '
         'corporation reimburses: rent 30%, home insurance, phone and internet (actual bills), '
         'and one Mercedes service. Rent and insurance carry no HST. Mercedes costs (lease, '
         'fuel, maintenance) are claimed at 80% business use, and the lease is also capped '
         'under s.235.'), F),
        (('Every file in this folder is losslessly compressed: images are pixel-identical to '
         'the originals, PDFs keep every page.'), F),
        (('Blue cells are inputs. Green cells pull from another sheet. Yellow cells are the '
         'judgement calls. Change a Yes/No in ITCs column F or an ITC % in column I and the '
         'return updates.'), F),
        ('', F),
        (('Open points are on the Questions sheet. Missing documents are marked NEEDED in the '
         'ITCs sheet Receipt column (filter it).'), FB),
    ]
    for r, (t, f) in enumerate(text, 1):
        c = ws.cell(r, 1, t)
        c.font = f
        c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 120


def sheet_questions(wb, need):
    ws = wb.create_sheet('Questions')
    ws['A1'], ws['A1'].font = 'Questions for you, and notes', FT
    ask = [
        (('Car lease HST paid in the period is 2,789.85 (776.55 at signing + 10 × 201.33). It is '
         'claimed at 80% business use × the 44.93% s.67.3 / s.235 ratio (99,500 list price) = '
         '35.94%. Using the 109,929.95 total list the ratio would be 40.67%. Agree with the cap '
         'and the list price? Lease cap sheet B12 and B22 drive it.'),
         'Lease ITC 1,002.69 (2,231.88 at 80% without the cap)'),
        (('Vehicle business use is set at 80% for the Mercedes (lease, fuel, maintenance) on '
         "Umar's estimate. There is no mileage log for it (odometer 14,035 km on Aug 12, 2026; "
         "in service Sep 30, 2025). Umar's 2025 mileage logs on his own car averaged about 700 "
         'business km a month. Is 80% supportable? Also, does a s.173 amount for the '
         'shareholder standby charge / operating benefit (Oct to Dec 2025) belong on line 104?'),
         'Line 104 is 0 now'),
        (('Mileage paid to Umar on his own car, Aug 15 to Sep 30, 2025 (Personal CC Expenses '
         'folder): claim a notional ITC under s.174?'), 'Small, not claimed'),
        ((f'Receipts missing for {need["n_missing"]} card purchases of $30 or more: '
         f'{need["missing_list"]}. Claimed on the card statement alone; acceptable?'),
         f'ITC {need["missing_itc"]:,.2f}'),
        (('Valet Car Wash AutoBilling receipts (04b, Apr to Aug) show "Ontario Tax" 5.20 but no '
         'GST/HST registration number. Claimed at 80%; acceptable?'), 'ITC 20.80'),
        (('Rogers ended April 2026: the Apr 22 bill was 163.58, then a 148.78 credit on the '
         'May 22 bill (repeated on Jun 22) that Rogers has not refunded. The credit is counted '
         'once, reducing the ITC. Agree?'), 'ITC reduced by 17.12'),
    ]
    notes = [
        ('Sales are counted by date received, as on your 2024-25 return. The Aug 22, 2025 deposit '
        'is in and the Aug 21, 2026 deposit is out; both are 8,542.80, so invoice-date basis '
        'would give the same result.'),
        ('iFly Calgary, Jul 11, 2026 (319.76): client entertainment per Umar, so claimed at 50% '
        'like meals (ITC 18.40).'),
        ('Credit Valley hospital, Jul 27, 2026 (1,200.00 + 21.00): personal, put on the corporate '
        'card by mistake. Not claimed; reduces the amount the corporation owes Umar. Halton '
        'Health Care parking (Oct 15, Nov 6, Nov 28, 2025; 78.25) is business and claimed.'),
        ('Mercedes service Apr 2, 2026 (353.94, invoice 1128781) was paid on Umar\'s personal '
        'Visa 1680; claimed as a director-funded expense at 80% (ITC 32.58). Receipts in '
        '04b/04-April.'),
        ('Petro-Canada Sep 9, 2025 (two charges): fuel for Umar\'s own car, for which mileage '
        'was paid until Sep 30. Not claimed.'),
    ]
    r = 3
    ws.cell(r, 1, 'Open questions').font = FB
    header(ws, r + 1, ['#', 'Question', 'Effect on the return'])
    r += 2
    for n, (q, eff) in enumerate(ask, 1):
        for c, v in enumerate((n, q, eff), 1):
            cell = ws.cell(r, c, v)
            cell.font = F
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    r += 1
    ws.cell(r, 1, 'Notes (settled)').font = FB
    r += 1
    for n, t in enumerate(notes, 1):
        for c, v in enumerate((n, t), 1):
            cell = ws.cell(r, c, v)
            cell.font = F
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        r += 1
    widths(ws, [4, 100, 26])


def main():
    receipts = copy_docs(write='--workpaper-only' not in sys.argv)
    t = pd.read_csv(LEDGER, parse_dates=['date'])
    p = t[t.gst_period == '2025-26'].copy()
    REFUNDED.update(refunded_purchases(p))
    wb = Workbook()
    sheet_readme(wb)
    wb.create_sheet('Return')
    _, sales_total = sheet_sales(wb, p)
    lease_ref = sheet_lease(wb)
    _, itc_total = sheet_itcs(wb, p, receipts, lease_ref)
    sheet_instalments(wb)
    card = p[(p.source == 'card') & (p.direction == 'out') & (p.amount >= 30)]
    missing = card[[key(x.date, x.description, x.amount) not in receipts
                    for x in card.itertuples()]]
    missing = missing[(missing.category != 'Personal — owed by Umar') & (missing.itc_claimable > 0)

                      & ~missing.apply(lambda x: key(x.date, x.description, x.amount)
                                       in REFUNDED, axis=1)]
    need = {'n_missing': len(missing), 'missing_itc': missing.itc_claimable.sum(),
            'missing_list': '; '.join(f'{x.description.split(" ")[0].title()} '
                                      f'{x.amount:,.2f} ({x.date:%b %-d, %Y})'
                                      for x in missing.itertuples()),
            'prof': p[p.category == 'Professional fees'].itc_claimable.sum()}
    sheet_questions(wb, need)
    sheet_return(wb, sales_total, itc_total)
    wb._sheets = [wb[n] for n in ['Read me', 'Return', 'Sales', 'ITCs', 'Lease cap',
                                  'Instalments', 'Questions']]
    path = OUT / '01 HST 2025-26 workpaper.xlsx'
    wb.save(path)
    print(path)


if __name__ == '__main__':
    main()
