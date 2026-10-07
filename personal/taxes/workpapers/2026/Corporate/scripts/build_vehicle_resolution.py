#!/usr/bin/env python3
"""
Write the sole director's resolution on personal use of the Mercedes.

    uv run --with python-docx python build_vehicle_resolution.py \
        --out "../../../../2026/Corporate/[Pending]Vehicle_Personal_Use_Resolution_2026.docx"

Schedule A figures come from `Vehicle personal use 2026.xlsx` (Summary tab). Update
SCHEDULE once the December statements replace the estimates, then re-run.
"""
import argparse

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

# (item, total cost, personal 20%), tax year 2026, actual Jan–Aug plus estimate Sep–Dec
SCHEDULE = [
    ('Lease payments (12 × $1,750.01)', 21000.12, 4200.02),
    ('Insurance', 8721.81, 1744.36),
    ('Fuel and car washes', 5521.55, 1104.31),
    ('Maintenance (Apr 2 service, paid by the director)', 353.94, 70.79),
]
ACTUAL_TO_AUG, ESTIMATE_SEP_DEC = 4728.35, 2391.13

RESOLUTIONS = [
    'The business use of the Vehicle is set at 80% and its personal use at 20%, effective '
    'January 1, 2026, and continuing until changed by a further resolution.',
    'The Director shall reimburse the Corporation for 20% of the lease payments, insurance, '
    'fuel, car washes, maintenance and repairs of the Vehicle that the Corporation pays in '
    'each calendar year.',
    'Where the Director pays a cost of the Vehicle personally, the Corporation shall reimburse '
    'the Director for 80% of that cost only.',
    'The reimbursement shall be settled by offset against the amount the Corporation owes the '
    'Director for business expenses paid on its behalf. The Corporation shall record the offset '
    'as a charge to the Director’s shareholder account dated no later than December 31 of '
    'each year. If that account cannot absorb the full amount, the Director shall pay the '
    'shortfall to the Corporation no later than 45 days after the end of the year.',
    'The amount for 2026 is calculated in Schedule A. The estimated months shall be replaced '
    'by actual amounts from the December 2026 statements before the year end entry is made.',
    'The Director shall keep a log of business and personal kilometres driven in the Vehicle, '
    'and the business use percentage shall be reviewed each year against that log.',
]


def money(x):
    return f'${x:,.2f}'


def para(doc, text, bold=False, size=11, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold, run.font.size = bold, Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    return p


def body(doc):
    centre = WD_ALIGN_PARAGRAPH.CENTER
    para(doc, '15248132 CANADA INC.', bold=True, size=14, align=centre, space_after=0)
    para(doc, 'operating as XYZ Bytes', size=10, align=centre, space_after=0)
    para(doc, '(the “Corporation”)', size=10, align=centre, space_after=14)
    para(doc, 'RESOLUTION OF THE SOLE DIRECTOR', bold=True, size=12, align=centre, space_after=0)
    para(doc, 'Personal use of the corporate vehicle', size=11, align=centre, space_after=14)
    para(doc, 'WHEREAS:', bold=True)
    recitals = [
        'The Corporation leases a 2026 Mercedes-Benz GLC 43 AMG, VIN W1NKM8HB1TF465706 '
        '(the “Vehicle”), from Mercedes-Benz Financial under a lease signed '
        'September 30, 2025, and pays its lease payments, insurance, fuel and upkeep;',
        'Umar Farooq Aslam, the sole director and shareholder (the “Director”), uses '
        'the Vehicle mainly for the business of the Corporation and also for personal driving;',
        'The Corporation owes the Director for business expenses the Director has paid '
        'personally on its behalf; and',
        'The personal use of the Vehicle should be paid for by the Director and not be a '
        'benefit conferred by the Corporation.',
    ]
    for i, text in enumerate(recitals):
        para(doc, f'{"ABCD"[i]}.  {text}')
    para(doc, 'NOW THEREFORE BE IT RESOLVED THAT:', bold=True)
    for i, text in enumerate(RESOLUTIONS, 1):
        para(doc, f'{i}.  {text}')
    para(doc, 'The foregoing resolution is consented to by the sole director of the Corporation '
              'in writing pursuant to the Canada Business Corporations Act.', space_after=30)
    para(doc, '______________________________', space_after=0)
    para(doc, 'Umar Farooq Aslam, Sole Director', space_after=12)
    para(doc, 'Dated: ____________________, 2026')


def schedule(doc):
    doc.add_page_break()
    para(doc, 'SCHEDULE A', bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, 'Personal use charge, tax year 2026', align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=12)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, ('Cost', 'Total cost', 'Personal 20%')):
        cell.text = ''
        cell.paragraphs[0].add_run(text).bold = True
    for item, total, personal in SCHEDULE:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = item, money(total), money(personal)
    row = table.add_row().cells
    row[0].text = ''
    row[0].paragraphs[0].add_run('Total charged to the Director').bold = True
    row[1].text = money(sum(t for _, t, _ in SCHEDULE))
    row[2].text = ''
    row[2].paragraphs[0].add_run(money(sum(p for *_, p in SCHEDULE))).bold = True
    for r in table.rows:
        for c in r.cells[1:]:
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    para(doc, '', space_after=6)
    para(doc, f'Actual, January to August 2026: {money(ACTUAL_TO_AUG)}. Estimated, September '
              f'to December 2026: {money(ESTIMATE_SEP_DEC)}. Insurance from October 1, 2026 is '
              'the renewed premium of $804.75 a month; fuel for the estimated months is the '
              'January to August monthly average.')
    para(doc, 'Of the total, $70.79 relates to a cost the Director paid personally. It is '
              'applied by reducing the Director’s claim for that cost; the effect on the '
              'shareholder account is the same.')
    para(doc, 'The 20% is charged on the full cost. The limit on deductible lease costs under '
              's.67.3 of the Income Tax Act applies only to the business 80%, and the portion '
              'above the limit remains a cost of the Corporation.')
    para(doc, 'Year end entry: debit Due to Shareholder, credit the vehicle expense accounts, '
              'by the final total.')


def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[1])
    ap.add_argument('--out', required=True)
    args = ap.parse_args()
    doc = Document()
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Times New Roman', Pt(11)
    body(doc)
    schedule(doc)
    doc.save(args.out)
    print(f'wrote {args.out}')


if __name__ == '__main__':
    main()
