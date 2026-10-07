#!/usr/bin/env python3
"""
Cross-check parsed bank statements against the workbook.

    python3 reconcile_debit_vs_workbook.py debit_2026_H1_parsed.csv \
            "/tmp/Business Transactions.xlsx" 2026-01-01 2026-06-30

Matches on (date, amount, direction). Description wording is deliberately NOT
matched -- the bank writes "MIRSERVICESPAY/PAY", the workbook writes "MIR pay".

Expect a few legitimate non-1:1 lines: the bank sometimes splits one logical
payment into two transfers, and the workbook sometimes splits one bank debit
into its components (e.g. a WISE payment plus its fee). Those show up here as
paired leftovers whose totals agree -- read them, don't "fix" them blindly.
"""
import csv
import datetime
import sys
from collections import Counter

import openpyxl

MASTER = 'T 23-26'


def num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return 0.0


def load_workbook_txns(path, lo, hi):
    M = openpyxl.load_workbook(path, data_only=True)[MASTER]
    out = []
    for r in M.iter_rows(min_row=1, max_row=1027, values_only=True):
        d = r[0]
        if not isinstance(d, datetime.datetime) or not (lo <= d.date() <= hi):
            continue
        payout, gross = num(r[6]), num(r[3])
        amount, direction = (payout, 'debit') if payout else (gross, 'credit')
        out.append({'date': d.date().isoformat(), 'desc': str(r[1]).strip(),
                    'amount': amount, 'dir': direction})
    return out


def main():
    csv_path, xlsx_path, lo, hi = sys.argv[1:5]
    lo = datetime.date.fromisoformat(lo)
    hi = datetime.date.fromisoformat(hi)

    with open(csv_path) as fh:
        stmt = [{**row, 'amount': float(row['amount'])} for row in csv.DictReader(fh)
                if lo <= datetime.date.fromisoformat(row['date']) <= hi]
    book = load_workbook_txns(xlsx_path, lo, hi)

    print(f'statement {len(stmt)} txns   workbook {len(book)} txns')

    key = lambda t: (t['date'], t['amount'], t['dir'])
    s, b = Counter(map(key, stmt)), Counter(map(key, book))
    only_s, only_b = s - b, b - s

    if not only_s and not only_b:
        print(f'EXACT MATCH on date + amount + direction, all {len(stmt)} lines')
    else:
        def show(diff, pool, label):
            print(f'\n{label}:')
            for k, count in diff.items():
                d = next(x for x in pool if key(x) == k)
                print(f'  {k[0]}  {k[2]:6} {k[1]:>11,.2f}  {d["desc"][:50]}  x{count}')
            return sum(k[1] * c for k, c in diff.items())

        ts = show(only_s, stmt, 'in statement, not in workbook')
        tb = show(only_b, book, 'in workbook, not in statement')
        print(f'\nunmatched totals: statement {ts:,.2f} vs workbook {tb:,.2f}'
              f'  -> {"same money, different granularity" if abs(ts - tb) < 0.011 else "REAL DIFFERENCE"}')

    # month-end balances
    print('\nbank closing balance by statement:')
    seen = {}
    for t in stmt:
        seen[t['file']] = (t['date'], float(t['balance']))
    for f, (d, bal) in sorted(seen.items(), key=lambda x: x[1][0]):
        print(f'  {f:26} {d}  {bal:>12,.2f}')


if __name__ == '__main__':
    main()
