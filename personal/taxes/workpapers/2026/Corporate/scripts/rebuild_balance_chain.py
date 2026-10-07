#!/usr/bin/env python3
"""
Rebuild the monthly starting-balance chain in Business Transactions.xlsx.

Every monthly tab holds its opening balance as a HARDCODED value in row 5.
Nothing enforces that it equals the previous month's row 3 total, so inserting,
deleting or re-dating any transaction silently leaves every later month stale.
This script recomputes the whole chain.

    # 1. recalculate, because the deltas must come from computed values
    soffice --headless --norestore --convert-to xlsx --outdir /tmp "Business Transactions.xlsx"
    # 2. rebuild, reading deltas from the recalculated copy, writing into the original
    python3 rebuild_balance_chain.py "Business Transactions.xlsx" "/tmp/Business Transactions.xlsx"
    # 3. recalculate again and verify
    soffice --headless --norestore --convert-to xlsx --outdir /tmp "Business Transactions.xlsx"
    python3 verify_workbook.py "/tmp/Business Transactions.xlsx"

Why two files: openpyxl cannot evaluate formulas. The recalculated copy supplies
the numbers; the original keeps its formulas and formatting and is what gets
written. Never save the recalculated copy over the original — LibreOffice
round-trips lose some styling.

Each month's delta (row 3 total minus row 5 opening) is independent of the
opening balance, so one pass is enough — no iteration needed.
"""
import shutil
import sys

import openpyxl

ORDER = ['Sep', 'Oct', 'Nov', 'Dec',
         'Jan', 'Feb', 'Mar', 'Apr', 'May', 'June', 'July', 'Aug 2024',
         'Sep 2024', 'Oct 2024', 'Nov 2024', 'Dec 2024',
         'Jan 2025', 'Feb 2025', 'Mar 2025', 'Apr 2025', 'May 2025', 'Jun 2025',
         'July 2025', 'Aug 2025', 'Sep 2025', 'Oct 2025', 'Nov 2025', 'Dec 2025',
         'Jan 2026', 'Feb 2026', 'Mar 2026', 'Apr 2026', 'May 2026', 'June 2026',
         'July 2026', 'Aug 2026']

COLS = (4, 5, 6, 7)      # D gross, E HST, F net, G payout
OPENING_ROW, TOTAL_ROW = 5, 3


def num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return 0.0


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    target, recalced = sys.argv[1], sys.argv[2]

    calc = openpyxl.load_workbook(recalced, data_only=True)
    if calc[ORDER[0]].cell(TOTAL_ROW, 4).value is None:
        sys.exit('the recalculated copy has no computed values — run LibreOffice on it first')

    deltas = {sh: [round(num(calc[sh].cell(TOTAL_ROW, c).value)
                         - num(calc[sh].cell(OPENING_ROW, c).value), 2) for c in COLS]
              for sh in ORDER}

    backup = target.replace('.xlsx', ' (pre-chain-rebuild).xlsx')
    shutil.copy(target, backup)
    print(f'backup written: {backup}\n')

    wb = openpyxl.load_workbook(target)
    running = [num(calc[ORDER[0]].cell(OPENING_ROW, c).value) for c in COLS]
    changed = 0

    print(f'{"month":12}{"opening D / E / F / G":>52}')
    for sh in ORDER:
        ws = wb[sh]
        before = [num(ws.cell(OPENING_ROW, c).value) for c in COLS]
        for i, c in enumerate(COLS):
            ws.cell(OPENING_ROW, c).value = round(running[i], 2)
        flag = ''
        if before != [round(x, 2) for x in running]:
            changed += 1
            flag = f'   was {before}'
        print(f'{sh:12}{str([round(x, 2) for x in running]):>52}{flag}')
        running = [round(a + b, 2) for a, b in zip(running, deltas[sh])]

    wb.calculation.fullCalcOnLoad = True
    wb.save(target)
    print(f'\nfinal closing: {running}')
    print(f'{changed} month(s) corrected · saved to {target}')
    print('now recalculate and run verify_workbook.py')


if __name__ == '__main__':
    main()
